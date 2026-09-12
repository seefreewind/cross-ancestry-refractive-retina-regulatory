#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_LINE_SPACING
from docx.shared import RGBColor
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]


def style_document(doc: Document, double_space: bool = True):
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        if name in styles:
            styles[name].font.name = "Times New Roman"
            styles[name].font.color.rgb = RGBColor(0, 0, 0)
    if "Title" in styles:
        styles["Title"].font.size = Pt(14)
        styles["Title"].font.bold = True
    if "Heading 1" in styles:
        styles["Heading 1"].font.size = Pt(12)
        styles["Heading 1"].font.bold = True
    if "Heading 2" in styles:
        styles["Heading 2"].font.size = Pt(12)
        styles["Heading 2"].font.bold = True
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    for p in doc.paragraphs:
        pf = p.paragraph_format
        if double_space:
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_after = Pt(0)


def add_paragraph(doc: Document, text: str, double_space: bool = True):
    p = doc.add_paragraph()
    p.add_run(text)
    if double_space:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    return p


def md_to_docx(md_path: Path, docx_path: Path, title_style: bool = True):
    doc = Document()
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            if title_style:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line[2:])
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(14)
            else:
                doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("|"):
            continue
        elif line.startswith("- ["):
            add_paragraph(doc, line)
        else:
            add_paragraph(doc, line)
    style_document(doc)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def add_tsv_table(doc: Document, path: Path, caption: str):
    doc.add_heading(caption, level=1)
    with path.open(encoding="utf-8") as f:
        rows = list(csv.reader(f, delimiter="\t"))
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    doc.add_paragraph()


def manuscript_docx():
    md = (ROOT / "manuscript/HUMAN_GENETICS_MANUSCRIPT_v1.2.md").read_text(encoding="utf-8")
    before_tables, after_tables = md.split("## Tables", 1)
    tables_part, after_legends = after_tables.split("## Figure legends", 1)
    legends_part, back_part = after_legends.split("## Data availability", 1)
    doc = Document()
    line_no = 1
    for raw in before_tables.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        else:
            p = add_paragraph(doc, f"{line_no:04d}  {line}")
            line_no += 1
    add_tsv_table(doc, ROOT / "submission/tables/Table1_GWAS_analysis_characteristics.tsv", "Table 1. GWAS and analysis characteristics")
    add_tsv_table(doc, ROOT / "submission/tables/Table2_main_cross_ancestry_architecture.tsv", "Table 2. Main cross-ancestry architecture")
    add_tsv_table(doc, ROOT / "submission/tables/Table3_DAR_heterogeneity_tests.tsv", "Table 3. DAR heterogeneity tests")
    doc.add_heading("Figure legends", level=1)
    for para in re.split(r"\n\s*\n", legends_part.strip()):
        add_paragraph(doc, para.replace("\n", " "))
    doc.add_heading("Data availability", level=1)
    for raw in ("## Data availability\n" + back_part).splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        else:
            add_paragraph(doc, line)
    style_document(doc)
    doc.save(ROOT / "submission/HUMAN_GENETICS_MANUSCRIPT_v1.2.docx")


def main():
    manuscript_docx()
    md_to_docx(ROOT / "submission/HUMAN_GENETICS_COVER_LETTER_v2.md", ROOT / "submission/HUMAN_GENETICS_COVER_LETTER_v2.docx", False)
    md_to_docx(ROOT / "supplement/HUMAN_GENETICS_SUPPLEMENT_v1.0.md", ROOT / "submission/HUMAN_GENETICS_SUPPLEMENT_v1.0.docx", False)
    md_to_docx(ROOT / "submission/STATEMENTS_AND_DECLARATIONS_v1.md", ROOT / "submission/STATEMENTS_AND_DECLARATIONS_v1.docx", False)
    final = ROOT / "submission/human_genetics_final"
    final.mkdir(exist_ok=True)
    for src, dst in [
        ("submission/HUMAN_GENETICS_MANUSCRIPT_v1.2.docx", "Manuscript.docx"),
        ("submission/HUMAN_GENETICS_COVER_LETTER_v2.docx", "Cover_Letter.docx"),
        ("submission/HUMAN_GENETICS_SUPPLEMENT_v1.0.docx", "Supplementary_Information.docx"),
        ("submission/STATEMENTS_AND_DECLARATIONS_v1.docx", "Statements_and_Declarations.docx"),
    ]:
        (final / dst).write_bytes((ROOT / src).read_bytes())


if __name__ == "__main__":
    main()
