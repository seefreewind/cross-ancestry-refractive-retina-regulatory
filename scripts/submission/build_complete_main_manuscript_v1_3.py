#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT_MD = ROOT / "manuscript" / "HUMAN_GENETICS_MANUSCRIPT_v1.3.md"
OUT_DOCX = ROOT / "submission" / "HUMAN_GENETICS_MAIN_MANUSCRIPT_COMPLETE_v1.3.docx"
FINAL_COPY = ROOT / "submission" / "human_genetics_final" / "Manuscript_complete_with_figures_tables_v1.3.docx"

TABLES = [
    (ROOT / "submission" / "tables" / "Table1_GWAS_analysis_characteristics.tsv", "Table 1. GWAS and analysis characteristics."),
    (ROOT / "submission" / "tables" / "Table2_main_cross_ancestry_architecture.tsv", "Table 2. Main cross-ancestry architecture results."),
    (ROOT / "submission" / "tables" / "Table3_DAR_heterogeneity_tests.tsv", "Table 3. DAR heterogeneity and robustness tests."),
]

FIGURES = [
    (
        ROOT / "figures" / "final_v2" / "Figure1_study_design_FINAL.png",
        "Fig. 1. Study design.",
    ),
    (
        ROOT / "figures" / "final_v2" / "Figure2_cross_ancestry_architecture_FINAL.png",
        "Fig. 2. Genome-wide and retinal-OCR cross-ancestry sharing.",
    ),
    (
        ROOT / "figures" / "final_v2" / "Figure3_DAR_heterogeneity_FINAL.png",
        "Fig. 3. DAR heterogeneity, permutation and precision.",
    ),
    (
        ROOT / "figures" / "final_v2" / "FigureS1_DAR_SLDXR_FINAL.png",
        "Supplementary Fig. S1. Descriptive ancestry-DAR S-LDXR estimate.",
    ),
]


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("^1^", "¹").replace("^2^", "²")
    text = text.replace("*Corresponding author:", "Corresponding author:")
    return text


def set_run_font(run, size: float = 12, bold: bool | None = None, italic: bool | None = None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body_paragraph(doc: Document, text: str, *, align=None, size: float = 12, bold: bool = False, italic: bool = False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0) if bold or align == WD_ALIGN_PARAGRAPH.CENTER else Inches(0.25)
    run = p.add_run(clean_inline(text))
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(clean_inline(text), level=level)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, size=12, bold=True)
    return p


def style_document(doc: Document):
    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in styles:
            styles[style_name].font.name = "Times New Roman"
            styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(12)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)


def parse_markdown_sections(text: str) -> dict[str, str]:
    matches = list(re.finditer(r"^## (.+)$", text, flags=re.MULTILINE))
    sections: dict[str, str] = {}
    title_block = text[: matches[0].start()] if matches else text
    sections["__title_block__"] = title_block.strip()
    for i, match in enumerate(matches):
        title = match.group(1).strip()
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections[title] = text[start:end].strip()
    return sections


def add_markdown_block(doc: Document, block: str, *, skip_heading: bool = False):
    for raw in block.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            add_body_paragraph(doc, line[2:], align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
        elif line.startswith("## "):
            if not skip_heading:
                add_heading(doc, line[3:], 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:], 2)
        else:
            add_body_paragraph(doc, line)


def add_title_block(doc: Document, block: str):
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if not lines:
        return
    title = lines[0][2:] if lines[0].startswith("# ") else lines[0]
    add_body_paragraph(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    for line in lines[1:]:
        italic = line.startswith("*Corresponding")
        add_body_paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=italic)


def add_tsv_table(doc: Document, path: Path, caption: str):
    add_body_paragraph(doc, caption, bold=True)
    with path.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f, delimiter="\t"))
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, row in enumerate(rows):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = clean_inline(val)
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                para.paragraph_format.space_after = Pt(0)
                if col_idx in {1, 2, 3, 4, 5, 6, 7, 8} and len(val) < 24:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_run_font(run, size=8.2, bold=(row_idx == 0))
    add_body_paragraph(doc, "")


def add_figures(doc: Document, legends: str):
    add_heading(doc, "Figure legends and embedded figures", 1)
    legend_paras = [p.replace("\n", " ").strip() for p in re.split(r"\n\s*\n", legends) if p.strip()]
    for i, (fig_path, short_caption) in enumerate(FIGURES):
        if i > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        matching = next((p for p in legend_paras if p.startswith(short_caption)), short_caption)
        add_body_paragraph(doc, matching, bold=True if matching == short_caption else False)
        if fig_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(fig_path), width=Inches(6.25))
        else:
            add_body_paragraph(doc, f"[Missing figure file: {fig_path}]", italic=True)


def build_docx():
    md = MANUSCRIPT_MD.read_text(encoding="utf-8")
    sections = parse_markdown_sections(md)
    doc = Document()
    style_document(doc)

    add_title_block(doc, sections["__title_block__"])

    ordered_main_sections = [
        "Abstract",
        "Keywords",
        "Introduction",
        "Materials and methods",
        "Results",
        "Discussion",
    ]
    for title in ordered_main_sections:
        if title in sections:
            add_heading(doc, title, 1)
            add_markdown_block(doc, sections[title], skip_heading=True)

    add_heading(doc, "Tables", 1)
    for table_path, caption in TABLES:
        add_tsv_table(doc, table_path, caption)

    if "Figure legends" in sections:
        add_figures(doc, sections["Figure legends"])

    for title in ["Data availability", "Code availability", "Acknowledgements", "Author contributions", "Funding", "Competing interests", "References"]:
        if title in sections:
            if title == "References":
                doc.add_section(WD_SECTION_START.NEW_PAGE)
            add_heading(doc, title, 1)
            add_markdown_block(doc, sections[title], skip_heading=True)

    style_document(doc)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    FINAL_COPY.parent.mkdir(parents=True, exist_ok=True)
    FINAL_COPY.write_bytes(OUT_DOCX.read_bytes())
    print(OUT_DOCX)
    print(FINAL_COPY)


if __name__ == "__main__":
    build_docx()
